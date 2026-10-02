"""Dumps the ball sheet's "Bowling Balls" and "staffMembersBag" tabs as JSON
for import-sheet.mjs. Needs openpyxl."""

import json
import sys
from datetime import date, datetime

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1])


def cell(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%b %Y')  # the sheet stores "Jul 2012" as a date
    if value is None:
        return ''
    if isinstance(value, str) and value.startswith('='):
        return ''  # the =IMAGE() preview column
    return value


rows = list(wb['Bowling Balls'].iter_rows(values_only=True))
balls = [[cell(v) for v in row[:19]] for row in rows[1:] if row[2]]

staff = []
for row in wb['staffMembersBag'].iter_rows(min_row=2, values_only=True):
    name = (row[0] or '').strip() if isinstance(row[0], str) else ''
    picks = [str(v).strip() for v in row[1:10] if v and str(v).strip()]
    if name:
        staff.append({'staff_member': name, 'sort_order': len(staff) + 1, 'balls': picks})

json.dump({'balls': balls, 'staff': staff}, sys.stdout)
